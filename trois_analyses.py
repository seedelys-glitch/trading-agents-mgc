# -*- coding: utf-8 -*-
"""
Trois analyses TradingAgents sur le même ticker, à la date du jour,
pour voir si les runs tombent d'accord (le « consensus » du guide).

- Reprend le fournisseur et les modèles choisis lors du dernier lancement
  de l'outil complet (raccourci « TradingAgents »).
- Force 1 seul tour de débat et 1 seul tour de discussion risque.
- Estime le nombre d'appels au modèle et attend le feu vert avant de lancer.
- Sauvegarde le rapport de chaque run dans son propre dossier.
- Affiche un tableau récap : run, verdict, argument baissier en une phrase.
- Ajoute une ligne au « Journal paper trading.xlsx ».
"""

import csv
import json
import os
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

DOSSIER = Path(__file__).resolve().parent
os.chdir(DOSSIER)  # pour que le fichier .env de ce dossier soit lu

from rich.console import Console  # noqa: E402
from rich.table import Table  # noqa: E402

console = Console()

NB_RUNS = 3
JOURNAL = DOSSIER / "Journal paper trading.xlsx"
PREFS = Path.home() / ".tradingagents" / "cli_prefs.json"

VERDICTS_FR = {
    "Buy": "Acheter",
    "Overweight": "Renforcer",
    "Hold": "Conserver",
    "Underweight": "Alléger",
    "Sell": "Vendre",
}


def stop(message):
    """Affiche un message et quitte proprement."""
    console.print(f"\n[bold red]{message}[/bold red]")
    sys.exit(1)


def texte_du_message(msg):
    """Récupère le texte d'une réponse de modèle, quel que soit son format."""
    contenu = getattr(msg, "content", msg)
    if isinstance(contenu, str):
        return contenu
    if isinstance(contenu, list):
        morceaux = []
        for bloc in contenu:
            if isinstance(bloc, str):
                morceaux.append(bloc)
            elif isinstance(bloc, dict) and bloc.get("type") == "text":
                morceaux.append(bloc.get("text", ""))
        return "\n".join(morceaux)
    return str(contenu)


def premiere_phrase(texte, limite=220):
    """Plan B si le résumé par le modèle échoue : début du texte."""
    texte = " ".join((texte or "").split())
    return (texte[:limite] + "…") if len(texte) > limite else (texte or "—")


def main():
    console.print("\n[bold]TradingAgents — 3 analyses pour tester le consensus[/bold]\n")

    # 1. Réglages du dernier lancement de l'outil complet
    try:
        prefs = json.loads(PREFS.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        prefs = {}
    if not prefs.get("llm_provider") or not prefs.get("quick_think_llm"):
        stop(
            "Lance d'abord l'outil complet une fois (raccourci « TradingAgents » sur le Bureau) :\n"
            "c'est là que tu choisis ton fournisseur d'IA et tes modèles légers.\n"
            "Ce script reprendra ensuite les mêmes choix."
        )

    # Imports lourds seulement après les vérifications rapides
    from cli.models import AnalystType, AssetType
    from cli.prompts import detect_asset_type, filter_analysts_for_asset_type, normalize_ticker_symbol
    from tradingagents.default_config import DEFAULT_CONFIG
    from tradingagents.graph.trading_graph import TradingAgentsGraph
    from tradingagents.llm_clients.api_key_env import get_api_key_env

    fournisseur = prefs["llm_provider"].lower()
    cle_env = get_api_key_env(fournisseur)
    if cle_env and fournisseur not in ("ollama", "openai_compatible") and not os.environ.get(cle_env):
        stop(
            f"Ta clé {cle_env} n'est pas dans le fichier .env.\n"
            "Lance l'outil complet (raccourci « TradingAgents ») : il te la demande en saisie masquée\n"
            "et la range lui-même dans .env."
        )

    # 2. Le ticker
    saisie = input("Ticker à analyser (ex : NVDA, MC.PA pour LVMH, XAUUSD pour l'or) : ").strip()
    if not saisie:
        stop("Aucun ticker saisi.")
    ticker = normalize_ticker_symbol(saisie)
    type_actif = detect_asset_type(ticker)
    jour = date.today().strftime("%Y-%m-%d")

    analystes_prefs = [a for a in prefs.get("analysts") or [] if a in {x.value for x in AnalystType}]
    if not analystes_prefs:
        analystes_prefs = [a.value for a in AnalystType]
    analystes = [
        a.value
        for a in filter_analysts_for_asset_type([AnalystType(a) for a in analystes_prefs], AssetType(type_actif))
    ]

    # 3. Configuration : mêmes modèles que le dernier lancement, 1 seul tour de débat
    config = DEFAULT_CONFIG.copy()
    config["llm_provider"] = fournisseur
    config["quick_think_llm"] = prefs["quick_think_llm"]
    config["deep_think_llm"] = prefs.get("deep_think_llm") or prefs["quick_think_llm"]
    if prefs.get("backend_url"):
        config["backend_url"] = prefs["backend_url"]
    config["output_language"] = os.environ.get("TRADINGAGENTS_OUTPUT_LANGUAGE") or prefs.get("output_language") or "French"
    config["max_debate_rounds"] = 1
    config["max_risk_discuss_rounds"] = 1

    # 4. Estimation des appels, puis feu vert
    n = len(analystes)
    bas, haut = n * 2 + 8, n * 5 + 8  # analystes (2 à 5 appels chacun) + 8 agents de débat/décision
    total_bas, total_haut = bas * NB_RUNS + 1, haut * NB_RUNS + 1  # +1 pour résumer les arguments baissiers

    console.print(f"\n[bold]Ticker[/bold] : {ticker} ({type_actif})   [bold]Date[/bold] : {jour}")
    console.print(f"[bold]Fournisseur[/bold] : {fournisseur}")
    console.print(f"[bold]Modèle quick[/bold] (le plus utilisé) : {config['quick_think_llm']}")
    console.print(f"[bold]Modèle deep[/bold] (les 2 managers) : {config['deep_think_llm']}")
    console.print(f"[bold]Analystes[/bold] : {', '.join(analystes)}   [bold]Débat[/bold] : 1 tour")
    console.print(
        f"\n[bold yellow]Estimation : environ {bas} à {haut} appels au modèle par analyse, "
        f"soit {total_bas} à {total_haut} appels pour les {NB_RUNS} analyses.[/bold yellow]"
    )
    console.print("[dim]Chaque appel est facturé par ton fournisseur. Regarde ta page de consommation après.[/dim]")

    feu_vert = input("\nTape OUI pour lancer les 3 analyses (autre chose = annuler) : ").strip().upper()
    if feu_vert not in ("OUI", "O"):
        console.print("Annulé, rien n'a été lancé.")
        return

    # 5. Les 3 analyses
    horodatage = datetime.now().strftime("%H%M%S")
    dossier_rapports = DOSSIER / "rapports" / f"{ticker.replace('/', '-').replace('=', '-')}_{jour}_{horodatage}"
    resultats = []
    dernier_graphe = None
    for i in range(1, NB_RUNS + 1):
        try:
            graphe = TradingAgentsGraph(selected_analysts=analystes, debug=False, config=config)
            dernier_graphe = graphe
            with console.status(f"Analyse {i}/{NB_RUNS} en cours… (plusieurs minutes, ne ferme pas la fenêtre)"):
                etat, verdict = graphe.propagate(ticker, jour, asset_type=type_actif)
            chemin = dossier_rapports / f"run_{i}"
            graphe.save_reports(etat, ticker, chemin, html=True)
            baissier = (etat.get("investment_debate_state") or {}).get("bear_history", "")
            resultats.append({"run": i, "verdict": verdict, "baissier": baissier, "dossier": chemin})
            console.print(f"[green]✓ Analyse {i} terminée : {verdict} ({VERDICTS_FR.get(verdict, '?')})[/green]")
        except Exception as erreur:  # un run raté ne bloque pas les autres
            resultats.append({"run": i, "verdict": "ERREUR", "baissier": "", "dossier": None, "erreur": str(erreur)})
            console.print(f"[red]✗ Analyse {i} a échoué : {erreur}[/red]")

    reussis = [r for r in resultats if r["verdict"] not in ("ERREUR",)]
    if not reussis:
        stop("Aucune analyse n'a abouti. Montre ce message à Claude.")

    # 6. Résumé de l'argument baissier de chaque run en une phrase (1 seul appel)
    resumes, cle = {}, ""
    try:
        consigne = (
            "Voici l'argumentaire du chercheur baissier pour plusieurs analyses du même titre. "
            "Réponds en français, uniquement avec ces lignes :\n"
            + "".join(f"{r['run']}: <une phrase qui résume l'argument baissier du run {r['run']}>\n" for r in reussis)
            + "CLE: <l'argument baissier le plus important, toutes analyses confondues, en une phrase>\n\n"
            + "\n\n".join(f"=== RUN {r['run']} ===\n{r['baissier'][:6000]}" for r in reussis)
        )
        reponse = texte_du_message(dernier_graphe.quick_thinking_llm.invoke(consigne))
        for ligne in reponse.splitlines():
            tete, _, reste = ligne.partition(":")
            tete = tete.strip().strip("*#- ").upper()
            reste = reste.strip().strip("*").strip()
            if tete.isdigit():
                resumes[int(tete)] = reste
            elif tete in ("CLE", "CLÉ"):
                cle = reste
    except Exception as erreur:
        console.print(f"[dim]Résumé automatique indisponible ({erreur}) : début du texte affiché à la place.[/dim]")
    for r in resultats:
        if r["verdict"] == "ERREUR":
            r["resume"] = premiere_phrase(r.get("erreur", "—"))
        else:
            r["resume"] = resumes.get(r["run"]) or premiere_phrase(r["baissier"])
    if not cle:
        cle = reussis[0]["resume"]

    # 7. Consensus
    compte = Counter(r["verdict"] for r in reussis if r["verdict"] in VERDICTS_FR)
    if compte and compte.most_common(1)[0][1] >= 2:
        v, nb = compte.most_common(1)[0]
        consensus = f"{v} ({nb}/{len(resultats)})"
    else:
        consensus = "Pas d'accord"

    # 8. Tableau récap
    tableau = Table(title=f"{ticker} — {jour}")
    tableau.add_column("Run", justify="center")
    tableau.add_column("Verdict final")
    tableau.add_column("Argument baissier (une phrase)")
    for r in resultats:
        v = r["verdict"]
        tableau.add_row(str(r["run"]), f"{v} ({VERDICTS_FR[v]})" if v in VERDICTS_FR else v, r["resume"])
    console.print()
    console.print(tableau)
    console.print(f"\n[bold]Consensus[/bold] : {consensus}")
    console.print(f"[bold]Argument baissier clé[/bold] : {cle}")
    console.print(f"\n[dim]Rapports complets (ouvre complete_report.html) : {dossier_rapports}[/dim]")

    # 9. Ligne ajoutée au journal de paper trading
    ligne = [
        jour,
        ticker,
        " / ".join(r["verdict"] for r in resultats),
        consensus,
        None,  # évolution à 5 séances : à remplir
        None,  # évolution à 1 mois : à remplir
        None,  # bon call ? : à remplir
        cle,
    ]
    try:
        from openpyxl import load_workbook

        classeur = load_workbook(JOURNAL)
        feuille = classeur["Journal"]
        rang = 2
        while feuille.cell(row=rang, column=1).value not in (None, ""):
            rang += 1
        for col, valeur in enumerate(ligne, start=1):
            if valeur is not None:
                feuille.cell(row=rang, column=col, value=valeur)
        classeur.save(JOURNAL)
        console.print(f"[green]✓ Ligne ajoutée au « {JOURNAL.name} » (ligne {rang}).[/green]")
    except PermissionError:
        secours = DOSSIER / "journal_a_recopier.csv"
        nouveau = not secours.exists()
        with secours.open("a", newline="", encoding="utf-8-sig") as f:
            ecrivain = csv.writer(f, delimiter=";")
            if nouveau:
                ecrivain.writerow(["Date", "Ticker", "Verdicts", "Consensus", "5 séances", "1 mois", "Bon call ?", "Argument baissier clé"])
            ecrivain.writerow(["" if v is None else v for v in ligne])
        console.print(
            f"[yellow]Le journal Excel est ouvert : la ligne a été mise dans « {secours.name} ». "
            "Ferme Excel la prochaine fois avant de lancer.[/yellow]"
        )
    except Exception as erreur:
        console.print(f"[yellow]Journal non mis à jour ({erreur}).[/yellow]")

    console.print(
        "\n[dim]Rappel du guide : un run seul est un tirage au sort. Lis l'argument baissier en premier, "
        "vérifie les chiffres à la source, et reste en paper trading.[/dim]"
    )


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        console.print("\nArrêté.")
